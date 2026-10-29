#!/usr/bin/env python3
"""Build matching A0 PDF and editable PowerPoint from a shared vector layout."""
from pathlib import Path
import json
import yaml
import argparse
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from html_poster import HTMLPoster

ROOT = Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--variant',choices=['source'],default='source')
parser.add_argument('--output-dir',default='output')
ARGS=parser.parse_args()
OUT=ROOT/ARGS.output_dir
VARIANT=ARGS.variant
OUT.mkdir(parents=True,exist_ok=True)
CFG = json.loads((ROOT / 'poster_config.json').read_text())
STYLE = yaml.safe_load((ROOT/'assets/poster.yaml').read_text())
TOKENS = STYLE['robopin']
W, H = TOKENS['width_mm'], TOKENS['height_mm']
MM = 72 / 25.4
NAVY = '#222222'
TEAL = '#E54C64'
MINT = '#E4F3F0'
INK = '#222222'
GRAY = '#566B7C'
BG = '#F0F4F7'
LINE = '#D5E0E7'
WHITE = '#FFFFFF'
BLUE = '#5782A3'
AMBER = '#B2711D'
font_dir = ROOT/'assets/fonts'
for name, file in [('Regular','OpenSans-Regular.ttf'),('Bold','OpenSans-Bold.ttf'),('Heading','ComicNeue-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
pdfmetrics.registerFont(TTFont('Mono','/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(W), Mm(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
c = canvas.Canvas(str(OUT / 'RoboPIN_A0_portrait.pdf'), pagesize=(W*MM,H*MM), initialFontName='Regular')
c.setTitle('RoboPIN — Grounded Embodied Reasoning via Pinned Chain-of-Thought')
c.setAuthor('Yaoting Huang et al. — Tianjin University')
c.setSubject('ACM Multimedia 2026 | A0 portrait | official logos')
audit = []
html = HTMLPoster(W,H)

def rgb(v): return RGBColor.from_string(v.lstrip('#'))

def rect(x,y,w,h,fill=WHITE,stroke=None,radius=0,lw=1):
    html.rect(x,y,w,h,fill,stroke,radius,lw)
    c.setFillColor(HexColor(fill))
    c.setLineWidth(lw)
    if stroke: c.setStrokeColor(HexColor(stroke))
    if radius: c.roundRect(x*MM,(H-y-h)*MM,w*MM,h*MM,radius*MM,fill=1,stroke=bool(stroke))
    else: c.rect(x*MM,(H-y-h)*MM,w*MM,h*MM,fill=1,stroke=bool(stroke))
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Mm(x),Mm(y),Mm(w),Mm(h))
    s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    if stroke: s.line.color.rgb=rgb(stroke); s.line.width=Pt(lw)
    else: s.line.fill.background()
    if radius:
        s.adjustments[0]=min(radius/min(w,h),0.15)
    return s

def line(x1,y1,x2,y2,color=LINE,lw=1.5):
    html.line(x1,y1,x2,y2,color,lw)
    c.setStrokeColor(HexColor(color)); c.setLineWidth(lw)
    c.line(x1*MM,(H-y1)*MM,x2*MM,(H-y2)*MM)
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Mm(x1),Mm(y1),Mm(x2),Mm(y2))
    s.line.color.rgb=rgb(color); s.line.width=Pt(lw)

def text(x,y,w,value,size=28,bold=False,color=INK,leading=1.22,mono=False,align='left',max_h=None,heading=False,highlights=()):
    font='Heading' if heading else ('Mono' if mono else ('Bold' if bold else 'Regular'))
    lines=[]
    for para in value.split('\n'):
        if not para:
            lines.append('');continue
        words=para.split(' '); row=''
        for word in words:
            test=(row+' '+word).strip()
            if row and pdfmetrics.stringWidth(test,font,size)>w*MM:
                lines.append(row);row=word
            else:row=test
        lines.append(row)
    height=len(lines)*size*leading/MM
    if max_h is not None and height>max_h:
        raise ValueError(f'Text overflows box: {value[:50]} ({height:.1f}>{max_h} mm)')
    import re
    def runs(t):
        if not highlights:return [(t,color)]
        pattern='('+'|'.join(re.escape(h) for h in sorted(highlights,key=len,reverse=True))+')'
        return [(part,TEAL if part in highlights else color) for part in re.split(pattern,t) if part]
    for i,t in enumerate(lines):
        tw=pdfmetrics.stringWidth(t,font,size)
        if tw>w*MM+0.1:raise ValueError(f'Text too wide: {t}')
        tx=x*MM + ((w*MM-tw)/2 if align=='center' else (w*MM-tw if align=='right' else 0))
        for part,part_color in runs(t):
            c.setFillColor(HexColor(part_color));c.setFont(font,size)
            c.drawString(tx,(H-y)*MM-size*0.86-i*size*leading,part)
            tx+=pdfmetrics.stringWidth(part,font,size)
    tb=slide.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(height+4))
    tf=tb.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    from pptx.enum.text import PP_ALIGN
    for i,t in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.font.name='Comic Neue' if heading else ('DejaVu Sans Mono' if mono else 'Open Sans')
        for part,part_color in runs(t):
            r=p.add_run();r.text=part;r.font.color.rgb=rgb(part_color)
        p.font.size=Pt(size);p.font.bold=bold or heading;p.font.color.rgb=rgb(color)
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*leading)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
    audit.append({'text':value,'x':x,'y':y,'width':w,'height':height,'size_pt':size})
    html.text(x,y,w,lines,size,bold or heading,color,leading,mono,align,heading=heading,colored_lines=[runs(t) for t in lines])
    return height

def image(file,x,y,w,h):
    html.image(file,x,y,w,h)
    c.drawImage(str(file),x*MM,(H-y-h)*MM,w*MM,h*MM,mask='auto')
    slide.shapes.add_picture(str(file),Mm(x),Mm(y),Mm(w),Mm(h))

def circle(x,y,r,fill=TEAL,stroke=WHITE):
    html.circle(x,y,r,fill,stroke)
    c.setFillColor(HexColor(fill));c.setStrokeColor(HexColor(stroke));c.setLineWidth(3)
    c.circle(x*MM,(H-y)*MM,r*MM,fill=1,stroke=1)
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL,Mm(x-r),Mm(y-r),Mm(r*2),Mm(r*2))
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill);s.line.color.rgb=rgb(stroke);s.line.width=Pt(3)

def arrow(x1,y,x2,color=TEAL):
    line(x1,y,x2,y,color,3)
    line(x2-4,y-3,x2,y,color,3);line(x2-4,y+3,x2,y,color,3)

def logo_placeholder(key,x,y,w,h):
    file=CFG.get(key+'_file')
    if file and (ROOT/file).is_file():
        from PIL import Image
        with Image.open(ROOT/file) as im:iw,ih=im.size
        scale=min(w/iw,h/ih)
        image(ROOT/file,x+(w-iw*scale)/2,y+(h-ih*scale)/2,iw*scale,ih*scale)
    else:
        rect(x,y,w,h,'#1C3B56','#55748B',2)
        text(x+3,y+9,w-6,CFG[key+'_label'],18,True,'#D8E5ED',align='center',max_h=h-10)

def panel(x,y,w,h,title,kind,size=34):
    rect(x,y,w,29,TEAL)
    text(x+6,y+3,w-12,title,46 if w>500 else 42,color=WHITE,heading=True,align='center',max_h=25)

rect(0,0,W,H,WHITE)
# Plain white canvas follows the supplied SDO typography/layout reference.
text(22,12,568,'RoboPIN: Grounded Embodied Reasoning\nvia Pinned Chain-of-Thought',67,color=TEAL,heading=True,leading=1.16,max_h=60)
text(24,76,568,'Yaoting Huang* · Yifu Yuan*† · Linqi Han · Chengwen Li',27,True)
text(24,90,568,'Shuoheng Zhang · Xianze Yao · Hongyao Tang · Yan Zheng · Jianye Hao†',25,True)
text(24,108,568,'Tianjin University · ACM Multimedia 2026 · Rio de Janeiro, Brazil',25)
text(24,125,568,'quark404@tju.edu.cn · yuanyf@tju.edu.cn',24)
text(24,138,568,'* Equal contribution   † Corresponding authors',18,color=GRAY)
logo_placeholder('mm_logo',596,16,227,47)
seal_y=73
from PIL import Image
with Image.open(ROOT/CFG['tju_logo_file']) as seal: seal_h=64*seal.height/seal.width
seal_top=seal_y+(64-seal_h)/2
c.saveState();mask=c.beginPath();mask.circle(679*MM,(H-seal_y-32)*MM,32*MM);c.clipPath(mask,stroke=0)
c.drawImage(str(ROOT/CFG['tju_logo_file']),647*MM,(H-seal_top-seal_h)*MM,64*MM,seal_h*MM,mask='auto');c.restoreState()
seal_pic=slide.shapes.add_picture(str(ROOT/CFG['tju_logo_file']),Mm(647),Mm(seal_top),Mm(64),Mm(seal_h))
seal_pic._pic.spPr.prstGeom.set('prst','ellipse')
html.image(ROOT/CFG['tju_logo_file'],647,seal_top,64,seal_h,round_image=True)
text(718,90,111,'Tianjin University',26,True,NAVY)

panel(14,153,414,288,'Motivation','',42)
image(ROOT/'assets/figure-1-cot-comparison.png',46,190,350,350*180.96/335.60)
text(22,386,398,
     '1. Text-based references can drift between similar objects. Coordinate-only grounding lacks explicit identity tracking.',
     28,True,leading=1.18,max_h=37,highlights=('drift','explicit identity tracking'))
text(22,416,398,
     '2. A model may correctly analyze object A, then silently point to object B. This reasoning-to-answer misalignment breaks reliable visual grounding.',
     28,True,leading=1.18,max_h=37,highlights=('object A','object B','reasoning-to-answer misalignment'))

panel(442,153,385,288,'Key Insight: Persistent Identity','',42)
for file,x,px,py in [('view-0.png',452,.634,.576),('view-1.png',665,.844,.706)]:
    text(x,191,148,'View 0' if x==452 else 'View 1',24,True,align='center')
    with Image.open(ROOT/'assets'/file) as im: vh=148*im.height/im.width
    image(ROOT/'assets'/file,x,207,148,vh)
    circle(x+148*px,207+vh*py,3.8,fill='#007E87')
    rect(x+5,212,69,13,'#007E87')
    text(x+7,213,65,'obj_01',23,True,WHITE,align='center')
arrow(617,263,648,color=INK)
text(452,332,365,'From isolated grounding to persistent anchors',29,True,align='center')
text(454,350,360,
     '• Spatial references are usually isolated annotations, without explicit identity tracking during reasoning.',
     28,True,leading=1.16,max_h=36,highlights=('isolated annotations','identity tracking'))
text(454,383,360,
     '• PinCoT binds each entity to a name, ID, view and point. Later steps reuse the same ID.',
     28,True,leading=1.16,max_h=36,highlights=('name, ID, view and point','same ID'))
text(454,416,360,
     '• Our insight: visual anchors act as persistent reasoning variables, making each step traceable to visual evidence.',
     28,True,leading=1.16,max_h=48,highlights=('Our insight:','persistent reasoning','variables'))
# Entire original Figure 2 artwork: upper data pipeline and lower CoT example.
panel(14,451,813,314,'Our Method: RoboPIN / PinCoT','',35)
method_asset=ROOT/'assets/figure-2-full-method.png'
with Image.open(method_asset) as im: method_h=700*im.height/im.width
image(method_asset,70.5,486,700,method_h)

panel(14,775,400,380,'Experimental Results','clipboard',33)
# Preserve the actual published plot, including its original legend and labels.
image(ROOT/'assets/performance-original.png',54,816,320,320*106.12/166.2)
text(26,1028,376,'Ablation studies',29,True,TEAL)
# Two compact, editable tables. Values come from Tables 5 and 6, not Table 7.
text(26,1044,179,'Reward design',25,True,NAVY)
text(222,1044,180,'Training stages',25,True,NAVY)
rect(26,1060,178,15,'#FFF0F3',radius=2)
rect(222,1060,180,15,'#FFF0F3',radius=2)
text(31,1063,112,'Setting',20,True)
text(148,1063,51,'EmbSpatial',17,True,align='right')
text(227,1063,70,'Stage',20,True)
text(299,1063,45,'Spatial',19,True,align='right')
text(351,1063,46,'Point',19,True,align='right')
for i,(name,value) in enumerate([('Full reward','84.7'),('w/o Pin','82.0'),('w/o Consistency','82.7')]):
    yy=1081+i*19
    text(31,yy,117,name,22,bold=i==0,color=TEAL if i==0 else INK)
    text(151,yy,48,value,24,bold=i==0,color=TEAL if i==0 else INK,align='right')
for i,(name,spatial,point) in enumerate([('SFT','57.0','66.0'),('+ CoT-SFT','62.0','63.6'),('+ RFT','67.7','67.9')]):
    yy=1081+i*19
    text(227,yy,70,name,21,bold=i==2,color=TEAL if i==2 else INK)
    text(299,yy,45,spatial,24,bold=i==2,color=TEAL if i==2 else INK,align='right')
    text(351,yy,46,point,24,bold=i==2,color=TEAL if i==2 else INK,align='right')
line(214,1047,214,1133,LINE,1)
text(26,1140,376,'Spatial reasoning: +10.7 points from SFT to RFT.',21,True,TEAL)

panel(422,775,405,380,'Qualitative Results','chart',34)
image(ROOT/'assets/robot-three-rows.png',449,816,350,350*231.81/497.2)
text(435,986,379,'Zero-shot xArm · successes / 20',24,True,TEAL)
for i,(label,val) in enumerate([('Spatial','20 / 20'),('Multi-view','14 / 20'),('Long-horizon','11 / 20')]):
    xx=435+i*128
    text(xx,1003,119,val,29,True,TEAL,align='center')
    text(xx,1020,119,label,21,align='center')
image(ROOT/'assets/figure-4-two-examples.png',493,1041,263,263*96.68/240.25)
rect(14,1167,813,18,TEAL)
text(20,1168,801,'Persistent Visual Anchors · Consistent Reasoning · Grounded Answers',29,True,WHITE,align='center')

for item in audit:
    if item['x']<0 or item['y']<0 or item['x']+item['width']>W+.01 or item['y']+item['height']>H:
        raise ValueError(f'Content outside page: {item}')
c.showPage();c.save()
slide.notes_slide.notes_text_frame.text=(
    'A0 portrait 841 × 1189 mm. SDO reference restyle: Comic Neue Bold headings, Open Sans body, coral bars, white background. HTML/CSS and native PowerPoint share the same layout. '
    'Method uses one unchanged crop of the full Figure 2 artwork: data construction above, the original CoT reasoning example below. Key Insight restores the earlier paired-view explanation. No QR code, QR placeholder or paper/code label is included. '
    'Performance is the untouched plot cropped from paper Figure 1. Robot visualization includes all three rows of Figure 3. '
    'Reward ablation: Table 5, controlled EmbSpatial setting (84.7 / 82.0 / 82.7). '
    'Cumulative training: Table 6, Spatial 57.0 / 62.0 / 67.7, Point 66.0 / 63.6 / 67.9. '
    'The controlled ablation averages are distinct from the headline results; Point does not improve monotonically. '
    'No cross-step coverage panel. Text and ablation tables are editable; paper images and logos are replaceable pictures. '
    'Trial-count discrepancy in the source caption remains documented in SOURCES.md. No QR is included.'
)
prs.save(str(OUT/'RoboPIN_A0_editable.pptx'))
html.save(OUT/'RoboPIN_A0.html',font_dir)
(OUT/'layout-metadata.json').write_text(json.dumps({'variant':'original-full-method','method_asset':'assets/figure-2-full-method.png','method_region':[14,451,827,768],'chart_region':[14,775,414,1155],'qualitative_asset':'assets/figure-4-two-examples.png','performance_asset':'assets/performance-original.png','robot_asset':'assets/robot-three-rows.png'},indent=2))
(OUT/'layout-audit.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False))
print(f'Built {OUT}/RoboPIN_A0_portrait.pdf and RoboPIN_A0_editable.pptx ({len(audit)} text blocks).')
