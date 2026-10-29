#!/usr/bin/env python3
"""Check actual deliverables: dimensions, evidence text, fonts and editability."""
from pathlib import Path
from zipfile import ZipFile
import json
import hashlib
import fitz
import argparse
import re
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--output-dir',default='output')
args=parser.parse_args()
out=root/args.output_dir
meta=json.loads((out/'layout-metadata.json').read_text())
doc = fitz.open(out / 'RoboPIN_A0_portrait.pdf')
assert len(doc) == 1, 'Poster must be a single page'
page = doc[0]
size_mm = [page.rect.width*25.4/72, page.rect.height*25.4/72]
assert abs(size_mm[0]-841) < .01 and abs(size_mm[1]-1189) < .01
assert page.rotation == 0
pdf_text = page.get_text()
for item in ['RoboPIN', '84.7', '82.0', '82.7', '57.0', '62.0', '67.7', '66.0', '63.6', '67.9', '20 / 20', '14 / 20', '11 / 20', 'PLACEHOLDER']:
    assert item in pdf_text, f'Missing expected source value: {item}'
for removed in ['model parameters','PinCoT reasoning samples','evaluation benchmarks','CrossPoint vs. Mimo-7B']:
    assert removed not in pdf_text, f'Old metric card remains: {removed}'
assert not re.search(r'\b(?:Figure|Fig\.?|Tables?)\s*\d',pdf_text), 'Removed number captions have returned'
assert 'cross-step' not in pdf_text.lower()
for block in page.get_text('dict')['blocks']:
    for row in block.get('lines', []):
        for span in row['spans']:
            box=fitz.Rect(span['bbox'])
            assert box.x0 >= 0 and box.y0 >= 0 and box.x1 <= page.rect.width and box.y1 <= page.rect.height, span['text']
native_spans=[]
for block in page.get_text('dict')['blocks']:
    for row in block.get('lines', []):
        for span in row['spans']:
            if span['text'].strip(): native_spans.append(span)
for i,a in enumerate(native_spans):
    for b in native_spans[i+1:]:
        overlap=fitz.Rect(a['bbox']) & fitz.Rect(b['bbox'])
        # Adjacent colored PDF runs can differ by a few hundredths of a point
        # between the renderer's advance widths and MuPDF's font metrics.
        same_baseline_seam = overlap.width < .1 and abs(a['origin'][1]-b['origin'][1]) < .01
        assert overlap.is_empty or overlap.get_area()<1 or same_baseline_seam, f"Native text collision: {a['text']} / {b['text']}"
fonts=[]
for font in page.get_fonts():
    xref, ext, kind, base, *_ = font
    assert xref and doc.extract_font(xref)[3], f'Font is not embedded: {base}'
    fonts.append(base)
prs = Presentation(out / 'RoboPIN_A0_editable.pptx')
assert len(prs.slides) == 1
assert abs(prs.slide_width / 36000 - 841) < .01
assert abs(prs.slide_height / 36000 - 1189) < .01
slide = prs.slides[0]
boxes=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
native_text = '\n'.join(s.text for s in boxes)
for item in ['Motivation', 'Key Insight: Persistent Identity', 'Our Method: RoboPIN / PinCoT', 'Experimental Results', 'Qualitative Results', 'Spatial', 'SFT', 'CoT-SFT', 'RFT / GRPO', 'Reward design', 'Training stages', 'w/o Pin', 'w/o Consistency', '+ CoT-SFT', '+ RFT']:
    assert item in native_text, f'Expected editable content is missing: {item}'
method_text=page.get_text(clip=fitz.Rect(*(v*72/25.4 for v in meta['method_region'])))
for item in ['A. How PinCoT Reasons','B. Data Construction & Training','obj_01','[634,576]','[844,706]','Movement left']:
    assert item in method_text, f'Missing method evidence: {item}'
assert method_text.count('<obj')==1, 'Show one full anchor tag, then reuse its identity'
config = json.loads((root / 'poster_config.json').read_text())
embedded = {hashlib.sha256(s.image.blob).hexdigest() for s in slide.shapes if s.shape_type == 13}
paper_figures=['assets/figure-1-cot-comparison.png','assets/figure-2-data-pipeline.png','assets/view-0.png','assets/view-1.png',meta['qualitative_asset'],meta['performance_asset'],meta['robot_asset']]
for asset in paper_figures:
    assert hashlib.sha256((root/asset).read_bytes()).hexdigest() in embedded, f'Original paper figure missing: {asset}'
    from PIL import Image
    with Image.open(root/asset) as im:
        ratio=im.width/im.height
    matching=[s for s in slide.shapes if s.shape_type == 13 and hashlib.sha256(s.image.blob).hexdigest()==hashlib.sha256((root/asset).read_bytes()).hexdigest()]
    assert all(abs((s.width/s.height)/ratio-1)<.005 for s in matching), f'Figure stretched: {asset}'
    if asset in ['assets/view-0.png','assets/view-1.png']:
        assert len(matching)==1 and 451<=matching[0].top/36000<633, 'Worked example should appear only in Method'
logo_assets=[]
for key in ['mm_logo_file', 'tju_logo_file']:
    asset = config.get(key)
    if asset:
        assert hashlib.sha256((root/asset).read_bytes()).hexdigest() in embedded, f'Logo not embedded: {asset}'
        logo_assets.append(asset)
assert 'LOGO PLACEHOLDER' not in pdf_text, 'Outdated logo placeholder remains'
assert not any('LOGO PLACEHOLDER' in s.text for s in boxes)
header_pixel = page.get_pixmap(matrix=fitz.Matrix(.1,.1)).pixel(1,1)
assert min(header_pixel)>=230, 'Title corner should remain light for readability'
assert hashlib.sha256((root/'assets/rio-wave-background.png').read_bytes()).hexdigest() not in embedded, 'Superseded decorative background returned'
assert not any(s.text.strip() == 'Our insight' for s in boxes), 'Duplicate insight heading returned'
assert any('Comic Neue' == p.font.name for s in boxes for p in s.text_frame.paragraphs), 'Reference heading font missing'
for s in slide.shapes:
    assert s.left >= 0 and s.top >= 0 and s.left+s.width <= prs.slide_width+1 and s.top+s.height <= prs.slide_height+1, s.name
with ZipFile(out / 'RoboPIN_A0_editable.pptx') as z:
    assert z.testzip() is None
page.get_pixmap(matrix=fitz.Matrix(.62,.62)).save(out / 'RoboPIN_preview.png')
mm=72/25.4
for name, bounds in [
    ('detail-method.png', meta['method_region']),
    ('detail-results.png', (14,meta['chart_region'][1]-7,827,1160)),
    ('detail-original-figures.png', (14,146 if meta['variant']!='clean' else 168,827,meta['chart_region'][1]-7)),
    ('detail-performance.png', (14,meta['chart_region'][1]-7,414,1160)),
]:
    page.get_pixmap(matrix=fitz.Matrix(.75,.75),clip=fitz.Rect(*(v*mm for v in bounds))).save(out/name)
report={
    'pdf_pages': len(doc), 'page_mm': size_mm, 'orientation': 'portrait',
    'embedded_fonts': fonts, 'editable_text_boxes': len(boxes),
    'pptx_shapes': len(slide.shapes), 'source_value_checks': 'passed',
    'native_text_collisions': 'none',
    'page_bounds': 'passed', 'pptx_archive_integrity': 'passed',
    'visual_review': 'PDF preview inspected; PowerPoint not rendered in Office in this environment',
    'embedded_logo_assets': logo_assets,
    'embedded_original_paper_figures': paper_figures,
    'paper_figure_aspect_ratios': 'preserved',
    'performance_chart': 'Original paper crop; no redrawn bars',
    'robot_rows': 3,
    'method': 'Published CoT example with one full anchor tag; data construction and training separated',
    'ablation_source': 'Tables 5 and 6; editable table text',
    'aggregate_chart_values': 'Original Figure 1 image retained',
    'old_metric_cards': 'removed',
    'header_background': 'Plain white; coral flat section bars',
    'layout_variant': meta['variant'],
    'design': 'SDO reference typography: Comic Neue headings, Open Sans body, coral inline emphasis',
    'intentional_placeholders': ['paper/code QR'],
}
(out / 'validation.json').write_text(json.dumps(report, indent=2))
print(json.dumps(report, indent=2))
